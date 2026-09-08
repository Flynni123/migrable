"""Read Excel workbooks through the optional openpyxl dependency."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from migrable.core.schema import TableSchema
from migrable.sources.base import DataSource, DataSourceError, Row
from migrable.sources.inspection import schema_from_rows


class ExcelDataSource(DataSource):
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def _workbook(self):
        if not self.path.is_file():
            raise DataSourceError(f"Excel file not found: {self.path}")
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise DataSourceError("Excel support requires `pip install -e '.[excel]'`") from error
        try:
            return load_workbook(self.path, read_only=True, data_only=True)
        except Exception as error:
            raise DataSourceError(f"Could not open Excel file {self.path}: {error}") from error

    def list_datasets(self) -> list[str]:
        workbook = self._workbook()
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def _worksheet(self, dataset: str | None):
        workbook = self._workbook()
        if not dataset:
            dataset = workbook.sheetnames[0]
        if dataset not in workbook.sheetnames:
            workbook.close()
            raise DataSourceError(f"Worksheet '{dataset}' does not exist")
        return workbook, workbook[dataset]

    @staticmethod
    def _headers(row: tuple[object, ...]) -> list[str]:
        headers = [str(value).strip() if value is not None else "" for value in row]
        if not headers or not any(headers):
            raise DataSourceError("The first worksheet row must contain column names")
        if any(not header for header in headers) or len(headers) != len(set(headers)):
            raise DataSourceError("Worksheet headers must be non-empty and unique")
        return headers

    def get_schema(self, dataset: str | None = None) -> TableSchema:
        workbook, worksheet = self._worksheet(dataset)
        try:
            iterator = worksheet.iter_rows(values_only=True)
            headers = self._headers(next(iterator, ()))
            rows = [
                dict(zip(headers, values, strict=False))
                for _, values in zip(range(100), iterator, strict=False)
            ]
            return schema_from_rows(worksheet.title, headers, rows)
        finally:
            workbook.close()

    def read_rows(
        self, dataset: str | None = None, *, batch_size: int = 1_000
    ) -> Iterator[list[Row]]:
        if batch_size < 1:
            raise ValueError("batch_size must be positive")

        def batches() -> Iterator[list[Row]]:
            workbook, worksheet = self._worksheet(dataset)
            try:
                iterator = worksheet.iter_rows(values_only=True)
                headers = self._headers(next(iterator, ()))
                batch: list[Row] = []
                for values in iterator:
                    batch.append(dict(zip(headers, values, strict=False)))
                    if len(batch) >= batch_size:
                        yield batch
                        batch = []
                if batch:
                    yield batch
            finally:
                workbook.close()

        return batches()
